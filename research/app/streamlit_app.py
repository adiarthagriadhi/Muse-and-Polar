from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path
import tempfile
import zipfile

import numpy as np
import matplotlib.pyplot as plt
from matplotlib.patches import Circle
import streamlit as st
from docx import Document
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from muse2_pipeline.io import list_samples, load_sample
from muse2_pipeline.prv09 import compute_prv_09
from muse2_pipeline.maps import interpolate_brain_activity_idw, BRAIN_MAP_POSITIONS
from muse2_pipeline.segments import parse_marks, ruas_stats
from muse2_pipeline.eeg import CH

st.set_page_config(page_title="Muse 2 Research View", layout="wide")
st.title("Muse 2 — tampilan riset")
st.caption("Unggah folder trial (zip) atau pakai folder lokal. Bukan aplikasi klien / bukan HRV Polar.")


def draw_map(ax, scores, title):
    xs, ys, grid = interpolate_brain_activity_idw(scores)
    xx, yy = np.meshgrid(xs, ys)
    ax.contourf(xx, yy, grid, levels=7, cmap="RdYlBu_r")
    ax.contour(xx, yy, grid, levels=7, colors="k", linewidths=0.3, alpha=0.35)
    ax.add_patch(Circle((0, 0), 1.0, fill=False, lw=1.0))
    for ch, (x, y) in BRAIN_MAP_POSITIONS.items():
        ax.scatter(x, y, s=18, c="k")
        ax.text(x, y + 0.12, f"{ch}\n{scores[ch]:.1f}", ha="center", fontsize=7)
    ax.set_aspect("equal")
    ax.set_xlim(-1.2, 1.2)
    ax.set_ylim(-1.2, 1.2)
    ax.axis("off")
    ax.set_title(title, fontsize=9)


def laporan_docx(row, meta, edges, band, tab, cmp_, prv, png_bytes, tmax):
    doc = Document()
    doc.add_heading("Laporan riset Muse 2", level=1)
    doc.add_paragraph(f"Partisipan: {row['participant']}  |  Kondisi: {row['condition']}")
    doc.add_paragraph(f"Mulai: {row.get('start_iso') or meta.get('start_iso') or '-'}  |  Durasi: {tmax:.0f} s")
    doc.add_paragraph(f"Titik pengamatan: {edges}  |  Pita: {band}")
    doc.add_paragraph("Bukan diagnosis dan bukan HRV Polar. PRV = variabilitas pulsa dahi.")
    doc.add_heading("Denyut dan pulsa dahi (PPG2)", level=2)
    doc.add_paragraph(
        "Gelombang PPG mentah tidak ditampilkan di laporan. "
        "Yang dipakai: interval antar-puncak (IBI) dari saluran infrared PPG2, metode 09-lite."
    )
    if prv.get("ok"):
        doc.add_paragraph(
            f"Laju denyut {prv['hr']:.1f} /menit; "
            f"PRV-SDNN dahi {prv['prv_sdnn']:.0f} ms; "
            f"PRV-RMSSD dahi {prv['prv_rmssd']:.0f} ms "
            f"(n={prv['n']} interval)."
        )
    else:
        doc.add_paragraph("PRV tidak dihitung (interval pulsa terlalu sedikit).")
    ev = (meta.get("eeg_validation") or {}).get("status", "-")
    pv = (meta.get("ppg_validation") or {}).get("status", "-")
    doc.add_heading("Mutu berkas", level=2)
    doc.add_paragraph(f"Status EEG: {ev}. Status pulsa: {pv}.")
    doc.add_heading("Tampilan", level=2)
    if png_bytes:
        doc.add_picture(BytesIO(png_bytes), width=Inches(6.3))
    else:
        doc.add_paragraph("(Gambar tidak tersimpan.)")
    if len(tab):
        doc.add_heading("Aktivitas pita per ruas", level=2)
        table = doc.add_table(rows=1, cols=8)
        table.style = "Table Grid"
        hdr = ["Ruas", "Jendela", "TP9", "AF7", "AF8", "TP10", "Asimetri AF", "Bacaan"]
        for i, h in enumerate(hdr):
            table.rows[0].cells[i].text = h
        for _, r in tab.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(r["ruas"])
            cells[1].text = str(int(r["n_win"]))
            for j, c in enumerate(["TP9", "AF7", "AF8", "TP10"]):
                val = r[c]
                cells[2 + j].text = "-" if val != val else f"{val:.1f}"
            a = r["AF_asimetri"]
            cells[6].text = "-" if a != a else f"{a:.3f}"
            cells[7].text = str(r.get("arti", ""))
    if len(cmp_):
        doc.add_heading("Perbandingan ruas", level=2)
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for i, h in enumerate(["Banding", "Ukuran", "Median 1", "Median 2", "Selisih", "p"]):
            table.rows[0].cells[i].text = h
        for _, r in cmp_.iterrows():
            cells = table.add_row().cells
            cells[0].text = str(r.get("perbandingan", ""))
            cells[1].text = str(r.get("ukuran", ""))
            for j, k in enumerate(["median_1", "median_2", "selisih", "p"]):
                v = r.get(k)
                cells[2 + j].text = "-" if v != v else f"{float(v):.3f}"
    doc.add_paragraph("Empat kanal kering Muse. Bandingkan PRV hanya antar rekaman Muse.")
    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()


def unpack_zip(upload) -> Path:
    tmp = Path(tempfile.mkdtemp(prefix="muse2_"))
    zpath = tmp / "data.zip"
    zpath.write_bytes(upload.getbuffer())
    with zipfile.ZipFile(zpath) as zf:
        zf.extractall(tmp / "unzip")
    unzip = tmp / "unzip"
    hits = list(unzip.rglob("EEG_raw.csv"))
    if not hits:
        return unzip
    return hits[0].parents[2] if hits[0].parent.name in ("experiment", "baseline") else unzip


with st.sidebar:
    st.header("Data")
    mode = st.radio("Sumber", ["Folder lokal", "Unggah ZIP"])
    trials = None
    if mode == "Folder lokal":
        default = str(Path.home() / "muse2-research" / "data" / "raw" / "trials")
        path = st.text_input("Folder trials", default)
        p = Path(path)
        if p.exists():
            trials = p
        else:
            st.warning("Folder tidak ditemukan.")
    else:
        up = st.file_uploader("ZIP berisi Pxx_Trial/experiment/EEG_raw.csv", type=["zip"])
        if up is not None:
            trials = unpack_zip(up)

    processed = Path.home() / "muse2-research" / "data" / "processed"
    titik = st.text_input("Titik pengamatan (detik)", "60")
    band = st.radio("Pita", ["delta", "theta", "alpha", "beta"], index=2, horizontal=True)

if trials is None:
    st.info("Pilih folder trials atau unggah ZIP.")
    st.stop()

inv = list_samples(trials)
if inv.empty:
    st.error("Tidak ada pasangan EEG_raw.csv + PPG_raw.csv.")
    st.stop()

label = st.selectbox("Sampel", inv["label"].tolist())
row = inv[inv["label"] == label].iloc[0]
data = load_sample(row["folder"], processed if processed.exists() else None)

eeg_f, ppg, bp, meta = data["eeg_f"], data["ppg"], data["bp"], data["meta"]
tmax = float(eeg_f["time_s"].iloc[-1])
edges = parse_marks(titik, tmax)
tab, cmp_ = ruas_stats(bp, edges, band)
prv = compute_prv_09(ppg["PPG2"].to_numpy(float), ppg["time_s"].to_numpy(float))

t0 = st.slider("Jendela peta (detik)", 0, max(0, int(tmax) - 20), 0, 5)
t1v = min(t0 + 20, tmax)

c1, c2, c3 = st.columns(3)
c1.metric("Denyut sesi", f"{prv['hr']:.0f} /menit" if prv["ok"] else "n/a")
c2.metric("PRV-RMSSD dahi", f"{prv['prv_rmssd']:.0f} ms" if prv["ok"] else "n/a")
c3.metric("Durasi", f"{tmax:.0f} s")

fig = plt.figure(figsize=(13, 9), layout="constrained")
gs = fig.add_gridspec(4, 2, height_ratios=[1.35, 0.7, 1.0, 0.65], width_ratios=[1.45, 1.0])

ax1 = fig.add_subplot(gs[0, 0])
for c in CH:
    ax1.plot(eeg_f["time_s"], eeg_f[c], lw=0.4, label=c)
for m in edges[1:-1]:
    ax1.axvline(m, color="k", lw=1.1, ls="--")
ax1.axvspan(t0, t1v, color="C0", alpha=0.10)
ax1.set_ylim(-80, 80)
ax1.legend(ncol=4, fontsize=7)
ax1.set_ylabel("EEG (µV)")
ax1.grid(True, alpha=0.3)
ax1.set_title(row["participant"].replace("_", " "))

ax2 = fig.add_subplot(gs[1, 0], sharex=ax1)
tp = prv.get("tp", np.array([]))
if len(tp) >= 2:
    ax2.plot(tp[:-1], np.diff(tp) * 1000.0, lw=0.8, color="C1")
ax2.set_ylim(400, 1500)
ax2.set_ylabel("IBI dahi (ms)")
ax2.set_xlabel("Waktu (detik)")
for m in edges[1:-1]:
    ax2.axvline(m, color="k", lw=1.1, ls="--")
ax2.axvspan(t0, t1v, color="C0", alpha=0.10)
ax2.grid(True, alpha=0.3)

ax3 = fig.add_subplot(gs[:2, 1])
if bp is not None and len(bp):
    tmid = (bp["t_start"] + bp["t_end"]) / 2
    w = bp[(tmid >= t0) & (tmid < t1v)]
    if len(w) == 0:
        w = bp
    scores = {c: float(w[f"{c}_{band}"].median()) for c in CH}
    draw_map(ax3, scores, f"Peta {t0:.0f}–{t1v:.0f} s · {band}")
else:
    ax3.axis("off")

nseg = max(len(tab), 1)
inner = gs[2, :].subgridspec(1, nseg)
if len(tab):
    for i, r in tab.iterrows():
        ax = fig.add_subplot(inner[0, i])
        if r["n_win"] > 0:
            sc = {c: float(r[c]) for c in CH}
            draw_map(ax, sc, f"{r['ruas']}\n{r['arti']}")
        else:
            ax.axis("off")

from scipy.signal import welch
from muse2_pipeline.eeg import FS_EEG
psd = gs[3, :].subgridspec(1, 4)
for i, ch in enumerate(CH):
    ax = fig.add_subplot(psd[0, i])
    f, pxx = welch(eeg_f[ch].to_numpy(), fs=FS_EEG, nperseg=int(2 * FS_EEG))
    m = (f >= 1) & (f <= 40)
    ax.semilogy(f[m], pxx[m], lw=0.9)
    ax.set_title(ch, fontsize=8)
    ax.grid(True, alpha=0.3)

png_buf = BytesIO()
fig.savefig(png_buf, format="png", dpi=140, bbox_inches="tight")
pdf_buf = BytesIO()
fig.savefig(pdf_buf, format="pdf", bbox_inches="tight")
png_bytes = png_buf.getvalue()
st.pyplot(fig, clear_figure=True)

st.subheader("Aktivitas pita per ruas")
if len(tab):
    st.dataframe(tab[["ruas", "n_win", "TP9", "AF7", "AF8", "TP10", "AF_asimetri", "arti"]], hide_index=True)
st.subheader("Perbandingan ruas")
if len(cmp_):
    st.dataframe(cmp_, hide_index=True)
st.caption(
    "Pulsa dahi = saluran PPG2 (infrared). Di sini digambar sebagai IBI (jarak antar denyut), "
    "bukan gelombang mentah. PRV dahi bukan HRV sabuk dada."
)

base = f"{row['participant']}_{row['condition']}_{band}"

c1, c2, c3 = st.columns(3)
with c1:
    st.download_button("Unduh gambar (PNG)", data=png_buf.getvalue(), file_name=f"{base}.png", mime="image/png")
with c2:
    st.download_button("Unduh gambar (PDF)", data=pdf_buf.getvalue(), file_name=f"{base}_gambar.pdf", mime="application/pdf")
with c3:
    try:
        blob = laporan_docx(row, meta, edges, band, tab, cmp_, prv, png_bytes, tmax)
        st.download_button(
            "Unduh laporan Word",
            data=blob,
            file_name=f"{base}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except Exception as e:
        st.warning(f"Word gagal: {e}")
