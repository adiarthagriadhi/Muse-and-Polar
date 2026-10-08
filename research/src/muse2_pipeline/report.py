"""Exports for Research View v2: a ZIP of CSV tables and a Word report."""

from __future__ import annotations

import io
import json
import zipfile

import numpy as np
import pandas as pd

from .analysis import Session
from .eeg import BANDS
from .maps import interpolate_brain_activity_idw
from .montage import positions as montage_positions, regions as montage_regions


def _fmt(v, d=2):
    return "-" if v is None or (isinstance(v, float) and not np.isfinite(v)) else f"{v:.{d}f}"


def segment_table(res: dict) -> pd.DataFrame:
    rows = []
    for s in res["segments"]:
        r = {"ruas": f"{s['t0']:.0f}–{s['t1']:.0f} s", "nama": s["name"], "t0_s": s["t0"], "t1_s": s["t1"],
             "jendela": s["n_win"], f"jendela_bersih_{res['pair_label']}": s["n_pair"]}
        for c, v, z, n in zip(res["channels"], s["values"], s["z"], s["n_clean"]):
            r[f"{c}_{res['band']}_{res['metric']}"] = v
            r[f"{c}_z_vs_acuan"] = z
            r[f"{c}_jendela_bersih"] = n
        r[f"asimetri_{res['pair_label']}"] = s["asym"]
        r["acuan"] = bool(s["is_ref"])
        for src in ("ppg", "polar"):
            st = s.get(src) or {}
            r[f"{src}_hr_bpm"] = st.get("hr")
            r[f"{src}_rmssd_ms"] = st.get("rmssd")
            r[f"{src}_sdnn_ms"] = st.get("sdnn")
        rows.append(r)
    return pd.DataFrame(rows)


def export_zip(sess: Session, res: dict, settings: dict) -> bytes:
    spec = sess.spec(res["art_uv"])
    cols = {"t_start_s": spec["t_start"], "t_mid_s": spec["t_mid"]}
    bad = _bad(res, sess.nch)
    for c_i, c in enumerate(sess.ch):
        for b in BANDS:
            cols[f"{c}_{b}_uV2"] = spec["power"][b][:, c_i]
        cols[f"{c}_total_uV2"] = spec["total"][:, c_i]
        cols[f"{c}_subseg_bersih"] = spec["n_clean"][:, c_i]
        cols[f"{c}_artefak"] = bad[:, c_i]
    win = pd.DataFrame(cols)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ruas.csv", segment_table(res).to_csv(index=False))
        zf.writestr("uji_berpasangan.csv", pd.DataFrame(res["tests"]["pairs"]).to_csv(index=False))
        zf.writestr("uji_kruskal.csv", pd.DataFrame(res["tests"]["omnibus"]).to_csv(index=False))
        zf.writestr("jendela_dayapita.csv", win.to_csv(index=False))
        for name, ser in (("denyut_ppg.csv", sess.ppg), ("denyut_polar.csv", sess.polar)):
            if len(ser["t"]):
                zf.writestr(name, pd.DataFrame({"t_s": ser["t"], "ibi_ms": ser["ibi"] * 1000,
                                                "valid": ser["valid"]}).to_csv(index=False))
        zf.writestr("pengaturan.json", json.dumps({"sesi": str(sess.folder), **settings}, indent=2, ensure_ascii=False))
    return buf.getvalue()


def _bad(res: dict, nch: int) -> np.ndarray:
    """(n_windows, n_channels) artifact mask from the analysis result."""
    return np.array([[ch == "1" for ch in b] for b in res["bad"]], bool).reshape(-1, nch)


def _topo(ax, scores, vmin, vmax, title, fmt=lambda v: f"{v:.3g}", pos=None):
    """Layered head map: 12 thin filled RdYlBu_r levels with subtle contour lines (as the web view)."""
    from matplotlib.patches import Circle, Ellipse, Polygon

    pos = pos or montage_positions(list(scores))
    xs, ys, grid = interpolate_brain_activity_idw(scores, positions=pos, grid_size=140)
    xx, yy = np.meshgrid(xs, ys)
    if vmax - vmin < 1e-9:
        vmin, vmax = vmin - 0.5, vmax + 0.5
    levels = np.linspace(vmin, vmax, 13)
    ax.contourf(xx, yy, grid, levels=levels, cmap="RdYlBu_r", extend="both")
    ax.contour(xx, yy, grid, levels=levels, colors="k", linewidths=0.2, alpha=0.25)
    ax.add_patch(Circle((0, 0), 1.0, fill=False, lw=0.8, color="#333"))
    ax.add_patch(Polygon([(-0.1, 0.995), (0, 1.1), (0.1, 0.995)], closed=False, fill=False, lw=0.8, color="#333"))
    for sx in (-1, 1):
        ax.add_patch(Ellipse((sx * 1.03, 0), 0.1, 0.32, fill=False, lw=0.8, color="#333"))
    many = len(pos) > 6
    for ch, (x, y) in pos.items():
        ax.scatter(x, y, s=6 if many else 10, c="k", zorder=3)
        txt = ch if many else f"{ch}\n{fmt(scores[ch])}"
        ax.text(x, y + (0.06 if many else 0.13), txt, ha="center", va="bottom", fontsize=4 if many else 5.5, zorder=4)
    ax.set_xlim(-1.18, 1.18)
    ax.set_ylim(-1.18, 1.18)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title(title, fontsize=7)


def figure_png(sess: Session, res: dict, scale: str = "map") -> bytes:
    """Summary figure; ``scale`` = 'map' (own colour range per head map) or 'shared'."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    segs = res["segments"]
    colors = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]
    fig = plt.figure(figsize=(7.2, 6.4), layout="constrained")
    n = max(1, len(segs))
    gs = fig.add_gridspec(3, n, height_ratios=[1.1, 0.8, 1.0])
    ax = fig.add_subplot(gs[0, :])
    spec = sess.spec(res["art_uv"])
    vals = sess.window_values(spec, res["band"], res["metric"])
    colors += ["#e87ba4"]
    bad = _bad(res, sess.nch)
    clean_vals = np.where(bad, np.nan, vals) if len(bad) else vals
    if sess.nch <= 4:  # Muse: one line per channel
        groups = {c: [i] for i, c in enumerate(sess.ch)}
    else:  # many channels: one line per scalp region
        groups = {r: [sess.ch.index(c) for c in cs] for r, cs in montage_regions(sess.ch).items()} \
            or {"Rata-rata": list(range(sess.nch))}
    for k, (name, idx) in enumerate(list(groups.items())[:5]):
        with np.errstate(invalid="ignore"):
            y = np.nanmean(clean_vals[:, idx], axis=1)
        ax.plot(spec["t_mid"], y, lw=1, color=colors[k], label=name)
    if res["metric"] == "abs":
        ax.set_yscale("log")
    ax.set_ylabel(f"{res['band']} ({res['unit']})", fontsize=7)
    ax.legend(ncol=4, fontsize=6, frameon=False, loc="upper right")
    ax2 = fig.add_subplot(gs[1, :], sharex=ax)
    drawn = 0
    for ser, lab, col in ((sess.ppg, "PPG", "#2a78d6"), (sess.polar, sess.heart_name, "#eb6834")):
        if len(ser["t"]):
            y = np.where(ser["valid"], ser["ibi"] * 1000, np.nan)
            ax2.plot(ser["t"], y, lw=1, color=col, label=lab)
            drawn += 1
    ax2.set_ylabel("IBI (ms)", fontsize=7)
    ax2.set_xlabel("waktu (s)", fontsize=7)
    if drawn:
        ax2.legend(ncol=2, fontsize=6, frameon=False, loc="upper right")
    else:
        ax2.text(0.5, 0.5, "tanpa data jantung", transform=ax2.transAxes, ha="center", va="center",
                 fontsize=7, color="#7a7973")
    for a in (ax, ax2):
        a.tick_params(labelsize=6)
        a.grid(True, alpha=0.25, lw=0.5)
        for e in res["edges"][1:-1]:
            a.axvline(e, color="#52514e", lw=0.8, ls="--")
        for s in segs:
            if s["i"] == res["ref"]:
                a.axvspan(s["t0"], s["t1"], color="#8a8983", alpha=0.08, lw=0)
    # log10 for absolute power (as in the web view); labels show the real value
    tf = (lambda v: np.log10(v) if v and v > 0 else None) if res["metric"] == "abs" else (lambda v: v)
    lab = (lambda v: f"{10 ** v:.3g}") if res["metric"] == "abs" else (lambda v: f"{v:.3g}")
    allv = [tf(v) for s in segs for v in s["values"] if v is not None]
    allv = [v for v in allv if v is not None]
    vmin, vmax = (min(allv), max(allv)) if allv else (0, 1)
    for i, s in enumerate(segs[:n]):
        a = fig.add_subplot(gs[2, i])
        vals = [tf(v) if v is not None else None for v in s["values"]]
        if all(v is not None for v in vals):
            lo, hi = (min(vals), max(vals)) if scale == "map" else (vmin, vmax)
            _topo(a, dict(zip(sess.ch, vals)), lo, hi, f"{s['t0']:.0f}–{s['t1']:.0f} s" + (" · acuan" if s["is_ref"] else ""), lab,
                  pos=montage_positions(sess.ch))
        else:
            a.axis("off")
            a.set_title(f"{s['t0']:.0f}–{s['t1']:.0f} s · data kurang", fontsize=7)
    out = io.BytesIO()
    fig.savefig(out, format="png", dpi=160, facecolor="white")
    plt.close(fig)
    return out.getvalue()


def report_docx(sess: Session, res: dict, scale: str = "map") -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    ov = sess.overview()
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Laporan riset Muse 2", level=1)
    doc.add_paragraph(f"Sesi: {ov['name']}  ·  Folder: {ov['label']}  ·  Mulai: {ov['start'] or '-'}")
    if ov["note"]:
        doc.add_paragraph(f"Catatan: {ov['note']}")
    doc.add_paragraph(
        f"Durasi {ov['duration']:.0f} s  ·  Pita {res['band']} ({res['unit']})  ·  "
        f"Ruas acuan {res['ref'] + 1}  ·  Ambang artefak {res['art_uv']:.0f} µV  ·  "
        f"Titik: {', '.join(f'{e:.0f}' for e in res['edges'][1:-1]) or '-'}"
    )
    kanal = (f"4 kanal disesuaikan Muse ({', '.join(f'{m} = ' + (c if '+' not in c else 'rata-rata ' + c.replace('+', ' & ')) for m, c in sess.muse_map.items())}; "
             + (f"dirujuk ulang ke {sess.muse_ref} seperti Muse)" if sess.muse_ref else "referensi asli file, tanpa Fpz)") if sess.muse_map
             else f"{sess.nch} kanal ({', '.join(sess.ch)})")
    jantung = {"all": "PPG + EKG/Polar", "ppg": "hanya PPG", "ecg": "hanya EKG/Polar", "none": "tanpa data jantung"}[sess.heart]
    doc.add_paragraph(f"Kanal dianalisis: {kanal}  ·  Data jantung: {jantung}")
    doc.add_heading("Ringkasan sesi", level=2)
    q = ov["quality"]
    hp, hpo = ov["hrv"]["ppg"], ov["hrv"]["polar"]
    lines = [
        f"Mutu EEG: {q['eeg']['status']} ({q['eeg']['rate_hz']} Hz); jendela bersih per kanal: "
        + ", ".join(f"{c} {100 * (f or 0):.0f}%" for c, f in zip(sess.ch, res["clean_frac"])),
        f"Mutu PPG: {q['ppg']['status']}",
        f"Denyut dahi (PPG): {_fmt(hp['hr'], 0)} /menit · PRV-RMSSD {_fmt(hp['rmssd'], 0)} ms · "
        f"SDNN {_fmt(hp['sdnn'], 0)} ms · interval valid {100 * (hp['valid_frac'] or 0):.0f}%",
    ]
    if hpo:
        lines.append(f"{sess.heart_name}: {_fmt(hpo['hr'], 0)} /menit · HRV-RMSSD {_fmt(hpo['rmssd'], 0)} ms · "
                     f"SDNN {_fmt(hpo['sdnn'], 0)} ms · pNN50 {_fmt(hpo['pnn50'], 0)} %")
    if ov["sync_offset_ms"] is not None:
        lines.append(f"Sinkronisasi Muse−Polar: {ov['sync_offset_ms']:+.0f} ms (diterapkan ke data Polar)")
    for ln in lines:
        doc.add_paragraph(ln, style="List Bullet")
    doc.add_picture(io.BytesIO(figure_png(sess, res, scale)), width=Inches(6.4))
    doc.add_paragraph(
        "Peta kepala berlapis (12 pita, biru = rendah, merah = tinggi); "
        + ("skala warna per peta (kontras maksimum, bandingkan lewat angka)." if scale == "map"
           else "skala warna sama untuk semua ruas.")
    )

    doc.add_heading("Per ruas", level=2)
    segs = res["segments"]
    heart = sess.heart_name.split()[0]
    t = doc.add_table(rows=1, cols=1 + len(segs))
    t.style = "Table Grid"
    t.rows[0].cells[0].text = f"{res['band']} ({res['unit']})"
    for j, s in enumerate(segs):
        t.rows[0].cells[1 + j].text = f"Ruas {s['i'] + 1}: {s['t0']:.0f}–{s['t1']:.0f} s" + (f"\n{s['name']}" if s["name"] else "")
    rows = [(f"Jendela bersih {res['pair_label']}", [f"{s['n_pair']}/{s['n_win']}" for s in segs])]
    rows += [(c, [_fmt(s["values"][k], 1) for s in segs]) for k, c in enumerate(sess.ch)]
    if res["pair"]:
        rows.append((f"Asimetri {res['pair_label']}", [_fmt(s["asym"], 3) for s in segs]))
    rows += [("Denyut PPG", [_fmt((s["ppg"] or {}).get("hr"), 0) for s in segs]),
             ("PRV-RMSSD PPG", [_fmt((s["ppg"] or {}).get("rmssd"), 0) for s in segs]),
             (f"RMSSD {heart}", [_fmt((s["polar"] or {}).get("rmssd"), 0) for s in segs])]
    for name, vals_ in rows:  # one row per measure, one column per segment (fits any channel count)
        c = t.add_row().cells
        c[0].text = name
        for j, v in enumerate(vals_):
            c[1 + j].text = v

    doc.add_heading("Uji antar ruas", level=2)
    doc.add_paragraph(
        "Mann–Whitney dua sisi pada jendela 4 s yang tidak tumpang tindih; p dikoreksi Holm per ukuran; "
        "δ Cliff positif berarti ruas kedua cenderung lebih tinggi."
    )
    rows = res["tests"]["pairs"]
    if rows:
        t = doc.add_table(rows=1, cols=7)
        t.style = "Table Grid"
        for i, h in enumerate(["Ukuran", "Perbandingan", "n", "Median", "p", "p Holm", "δ Cliff"]):
            t.rows[0].cells[i].text = h
        for r in rows:
            c = t.add_row().cells
            c[0].text = r["measure"]
            c[1].text = r["pair"]
            c[2].text = f"{r['n_a']} / {r['n_b']}"
            c[3].text = f"{_fmt(r['med_a'], 3)} → {_fmt(r['med_b'], 3)}"
            c[4].text = _fmt(r["p"], 3)
            c[5].text = _fmt(r["p_holm"], 3) + (" *" if r["sig"] else "")
            c[6].text = f"{_fmt(r['delta'], 2)} ({r['effect']})"
    for o in res["tests"]["omnibus"]:
        doc.add_paragraph(f"Kruskal–Wallis {o['measure']}: p = {_fmt(o['p'], 3)} ({o['k']} ruas)")

    doc.add_heading("Catatan", level=2)
    doc.add_paragraph(
        "Nilai z adalah perubahan relatif terhadap ruas acuan, bukan diagnosis. "
        "Denyut PPG dihitung dengan pipeline B1 (lihat README). PRV dahi (PPG) tidak sama dengan HRV dari EKG/sabuk dada; "
        "pada B1, RMSSD ultra-singkat dari PPG dahi ditaksir berlebih dibanding Polar H10. "
        + (f"Empat elektroda kering Muse: peta kepala adalah interpolasi kasar." if sess.nch <= 4 else
           f"Peta kepala diinterpolasi dari {sess.nch} elektroda (posisi 10-20/10-10), gambaran kasar.")
    )
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
